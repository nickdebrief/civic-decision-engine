"""Stage 78B2B1 governed pathway visible-output equivalence tests."""

from __future__ import annotations

import copy
import hashlib
import html
import json
import sys
import tempfile
import types
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
ENGINE = ROOT / "scripts" / "evidence_led_governance_pipeline"
sys.path.insert(0, str(ENGINE))

import output_validation as validation  # noqa: E402
import report_adapter as adapter  # noqa: E402

W = validation.WORD_NAMESPACE


KIND_TO_SECTION = {
    "governed_observation": "Observations",
    "governed_inference": "Inferences",
    "governed_allegation": "Allegations",
    "governed_response": "Responses and express declinations",
    "governed_characterisation": "Characterisations",
    "decision_authority": "Authority and mandate",
    "decision_mandate": "Authority and mandate",
    "governed_determination": "Determinations and reasons",
    "determination_effect_event": "Determinations and reasons",
    "governed_challenge": "Challenges",
    "challenge_event": "Challenges",
    "challenge_outcome": "Challenges",
    "governed_remedy": "Remedies",
    "implementation_event": "Implementation, compliance and verification",
    "formal_completion_determination": "Implementation, compliance and verification",
    "determination_publication": "Determination publications",
    "procedural_notice": "Notice and participation",
    "procedural_deadline": "Procedural chronology",
    "procedural_time_event": "Procedural chronology",
    "deadline_calculation": "Procedural chronology",
    "pathway_link": "Pathway relationships",
}

SECTION_ORDER = [
    "Report identity and scope",
    "Authority and mandate",
    "Procedural chronology",
    "Notice and participation",
    "Observations",
    "Inferences",
    "Allegations",
    "Responses and express declinations",
    "Characterisations",
    "Determinations and reasons",
    "Challenges",
    "Remedies",
    "Implementation, compliance and verification",
    "Determination publications",
    "Pathway relationships",
    "Contested matters",
    "Scoped gaps",
    "Provenance and limitations",
]


def _row(kind: str, index: int) -> dict:
    identity = f"{kind}:logical-{index}"
    return {
        "object_kind": kind,
        "governed_logical_identity": identity,
        "parent_governed_identity": f"parent:{index}" if index % 3 == 0 else None,
        "endpoint_identities": [{"object_type": "canonical_record", "object_governed_identity": "canonical_record:CR-STAGE78", "relationship_role": "scope"}] if index % 4 == 0 else [],
        "category": f"category-alpha-{index}",
        "status": f"status-{index}",
        "source_authority_key": f"source-key-{index}",
        "row_authority_digest": hashlib.sha256(f"row:{identity}".encode()).hexdigest(),
        "ownership_path": f"canonical_record:CR-STAGE78 > {identity}",
        "represented_time": "2026-01-01" if index % 2 == 0 else None,
        "recorded_at": "2026-01-02T00:00:00Z",
        "chronology": {"basis": "represented_time" if index % 2 == 0 else "recording_time", "precision": "day" if index % 2 == 0 else "indeterminate", "lower_bound": "2026-01-01" if index % 2 == 0 else None, "upper_bound": "2026-01-01" if index % 2 == 0 else None, "ordering_relation": "ordered" if index % 2 == 0 else "indeterminate"},
        "source_bindings": [{"source_type": "synthetic", "source_id": f"SRC-{index}"}],
        "object_links": [{"object_type": "synthetic", "object_governed_identity": f"linked:{index}", "relationship_role": "context"}],
        "contestation": {"status": "not_represented"},
        "supersession": {"superseded": False, "withdrawn": False, "ceased": False},
        "reliance": {"status": "unavailable"},
        "limitations": [f"limitation-alpha-{index}"],
        "epistemic_label": f"epistemic-label-{index}",
        "attribution": {"mode": "synthetic"},
        "representation_mode": "synthetic",
        "review_state": {"status": "accepted"},
        "contrary_sources": [],
        "does_not_establish": {"observation_into_evidence": True, "inference_into_fact": True, "allegation_into_finding": True, "response_into_resolution": True, "authority_into_legal_validity": True, "determination_into_endorsement": True, "challenge_into_invalidation": True, "remedy_into_implementation": True, "implementation_report_into_verification": True, "compliance_evidence_into_proof": True, "verification_into_formal_completion": True, "publication_into_endorsement": True, "absence_into_non_occurrence": True},
    }


def _model() -> dict:
    rows = [_row(kind, index) for index, kind in enumerate(KIND_TO_SECTION, start=1)]
    sections = []
    for title in SECTION_ORDER:
        if title in {"Report identity and scope", "Provenance and limitations", "Scoped gaps"}:
            sections.append({"title": title, "rows": []})
            continue
        section_rows = [row for row in rows if KIND_TO_SECTION[row["object_kind"]] == title]
        if section_rows:
            sections.append({"title": title, "rows": section_rows})
    return {
        "render_model_contract": "stage78.pathway_render_model.v1",
        "report_type": "procedural_pathway_report",
        "specification_schema_version": "stage78.procedural_pathway_report_specification.v1",
        "title": "Synthetic pathway",
        "purpose": "Equivalence testing",
        "intended_audience": "Internal",
        "distribution_class": "internal_working",
        "canonical_record_reference": "CR-STAGE78",
        "projection_contract": "stage78.pathway_projection.v1",
        "projection_version": "78a2b2",
        "projection_digest": "a" * 64,
        "inclusion_mode": "full_canonical_record",
        "exclusion_rule": "No selected subset; full Canonical Record projection frozen.",
        "coverage": {"stage62_schema_present": True, "governed_observation": {"state": "available", "row_count": 1}},
        "unavailable_families": ["stage73"],
        "gaps": [{"statement": "No governed remedy record was found linked within this Canonical Record scope.", "binding_mechanism": "exact_governed_link"}],
        "sections": sections,
        "row_count": len(rows),
        "limitations": ["The procedural pathway report presents governed records only.", "Rendering does not establish legality, truth, fairness, completeness, compliance, publication or endorsement."],
    }


def _spec(model: dict) -> dict:
    return {
        "specification_schema_version": model["specification_schema_version"],
        "report_type": model["report_type"],
        "title": model["title"],
        "purpose": model["purpose"],
        "intended_audience": model["intended_audience"],
        "distribution_class": model["distribution_class"],
        "primary_record": {"reference": model["canonical_record_reference"]},
        "pathway_projection": {key: model[key] for key in ("canonical_record_reference", "projection_contract", "projection_version", "projection_digest", "inclusion_mode", "exclusion_rule", "coverage", "unavailable_families", "gaps")},
        "qualifications": ["Frozen qualification text"],
        "requested_formats": ["docx", "html"],
        "publication_engine_version": adapter.ENGINE_VERSION,
        "rendering_profile": "internal_pathway_v1",
        "template_version": "cde-internal-pathway-v1",
    }


def _qualification() -> dict:
    return {"review_mode": "four_gate", "disclosure_version": "stage75.v1", "disclosure": "Qualified for internal working review only.", "qualification_id": "QUAL-1", "qualification_digest": "b" * 64}


def _units(model: dict | None = None, spec: dict | None = None, qualification: dict | None = None) -> list[validation.PathwayVisibleUnit]:
    model = model or _model()
    spec = spec or _spec(model)
    return validation.canonical_pathway_units(model, spec, qualification)


def _write_docx(path: Path, units: list[validation.PathwayVisibleUnit], *, mutate: dict[int, str] | None = None) -> None:
    document = Document()
    for index, unit in enumerate(units):
        document.add_paragraph((mutate or {}).get(index, unit.value))
    document.save(path)


def _w_attr(name: str) -> str:
    return f'w:{name}'


def _escape_xml(value: str) -> str:
    return html.escape(value, quote=True)


def _paragraph_xml(
    text: str,
    *,
    bookmark: tuple[str, str] | None = None,
    anchor: tuple[str, str] | None = None,
    hidden: bool = False,
    deleted: bool = False,
    instruction: bool = False,
    wrapper: str | None = None,
    outside_after: str = "",
) -> str:
    starts = []
    ends = []
    if anchor:
        starts.append(f'<w:bookmarkStart w:id="{anchor[0]}" w:name="{anchor[1]}"/>')
        ends.insert(0, f'<w:bookmarkEnd w:id="{anchor[0]}"/>')
    if bookmark:
        starts.append(f'<w:bookmarkStart w:id="{bookmark[0]}" w:name="{bookmark[1]}"/>')
        ends.insert(0, f'<w:bookmarkEnd w:id="{bookmark[0]}"/>')
    props = "<w:rPr><w:vanish/></w:rPr>" if hidden else ""
    tag = "w:instrText" if instruction else "w:delText" if deleted else "w:t"
    run = f"<w:r>{props}<{tag}>{_escape_xml(text)}</{tag}></w:r>"
    if wrapper == "moveFrom":
        run = f"<w:moveFrom>{run}</w:moveFrom>"
    elif wrapper == "del":
        run = f"<w:del>{run}</w:del>"
    elif wrapper == "alternate":
        run = f"<mc:AlternateContent><mc:Choice>{run}</mc:Choice><mc:Fallback>{run}</mc:Fallback></mc:AlternateContent>"
    elif wrapper == "textbox":
        run = f"<w:drawing><w:txbxContent>{run}</w:txbxContent></w:drawing>"
    elif wrapper == "unsupported":
        run = f"<w:unsupportedTextWrapper>{run}</w:unsupportedTextWrapper>"
    outside = f"<w:r><w:t>{_escape_xml(outside_after)}</w:t></w:r>" if outside_after else ""
    return f"<w:p>{''.join(starts)}{run}{''.join(ends)}{outside}</w:p>"


def _write_framed_docx(path: Path, units: list[validation.PathwayVisibleUnit], *, authority_units: list[validation.PathwayVisibleUnit] | None = None, mutate_text: dict[int, str] | None = None, mutate_token: dict[int, str] | None = None, table_index: int | None = None, duplicate_name_index: int | None = None, duplicate_id_index: int | None = None, missing_end_index: int | None = None, hidden_index: int | None = None, deleted_index: int | None = None, instruction_index: int | None = None, metadata_only_index: int | None = None, unknown_extra: bool = False, spanning_pair: tuple[int, int] | None = None, overlapping_pair: tuple[int, int] | None = None) -> None:
    frames = validation.pathway_docx_frame_contract(authority_units or units)
    paragraphs = []
    next_id = 1
    duplicate_name = None
    duplicate_id = None
    if duplicate_name_index is not None:
        duplicate_name = frames["unit_token_by_frame"][units[0].framing_identity]
    if duplicate_id_index is not None:
        duplicate_id = "1"
    if spanning_pair is not None:
        first, second = spanning_pair
        token = frames["unit_token_by_frame"][units[first].framing_identity]
        paragraphs.append(f'<w:p><w:bookmarkStart w:id="900" w:name="{token}"/><w:r><w:t>{_escape_xml(units[first].value)}</w:t></w:r></w:p>')
        paragraphs.append(f'<w:p><w:r><w:t>{_escape_xml(units[second].value)}</w:t></w:r><w:bookmarkEnd w:id="900"/></w:p>')
        _write_docx_xml(path, paragraphs)
        return
    if overlapping_pair is not None:
        first, second = overlapping_pair
        token_a = frames["unit_token_by_frame"][units[first].framing_identity]
        token_b = frames["unit_token_by_frame"][units[second].framing_identity]
        paragraphs.append(f'<w:p><w:bookmarkStart w:id="901" w:name="{token_a}"/><w:bookmarkStart w:id="902" w:name="{token_b}"/><w:r><w:t>{_escape_xml(units[first].value)}</w:t></w:r><w:bookmarkEnd w:id="901"/><w:bookmarkEnd w:id="902"/></w:p>')
        _write_docx_xml(path, paragraphs)
        return
    for index, unit in enumerate(units):
        anchor = None
        if unit.unit_kind == "heading" and unit.field_identity in {"chapter", "section"}:
            anchor = (str(next_id), frames["structural_token_by_frame"][unit.framing_identity])
            next_id += 1
        if metadata_only_index == index:
            paragraphs.append(f'<w:p><w:bookmarkStart w:id="{next_id}" w:name="{frames["unit_token_by_frame"][unit.framing_identity]}"/><w:bookmarkEnd w:id="{next_id}"/></w:p>')
            next_id += 1
            continue
        token = (mutate_token or {}).get(index, frames["unit_token_by_frame"][unit.framing_identity])
        if duplicate_name_index == index and duplicate_name is not None:
            token = duplicate_name
        bookmark_id = duplicate_id if duplicate_id_index == index and duplicate_id is not None else str(next_id)
        next_id += 1
        text = (mutate_text or {}).get(index, unit.value)
        xml = _paragraph_xml(
            text,
            bookmark=(bookmark_id, token),
            anchor=anchor,
            hidden=hidden_index == index,
            deleted=deleted_index == index,
            instruction=instruction_index == index,
        )
        if missing_end_index == index:
            xml = xml.replace(f'<w:bookmarkEnd w:id="{bookmark_id}"/>', "")
        if table_index == index:
            xml = f"<w:tbl><w:tr><w:tc>{xml}</w:tc></w:tr></w:tbl>"
        paragraphs.append(xml)
    if unknown_extra:
        paragraphs.append(_paragraph_xml("unexpected", bookmark=(str(next_id), validation.PATHWAY_DOCX_VISIBLE_PREFIX + "f" * 35)))
    _write_docx_xml(path, paragraphs)


def _write_framed_docx_with_override(path: Path, units: list[validation.PathwayVisibleUnit], index_to_override: int, override_xml: str) -> None:
    frames = validation.pathway_docx_frame_contract(units)
    paragraphs = []
    next_id = 1
    for index, unit in enumerate(units):
        anchor = None
        if unit.unit_kind == "heading" and unit.field_identity in {"chapter", "section"}:
            anchor = (str(next_id), frames["structural_token_by_frame"][unit.framing_identity])
            next_id += 1
        bookmark = (str(next_id), frames["unit_token_by_frame"][unit.framing_identity])
        next_id += 1
        if index == index_to_override:
            paragraphs.append(override_xml.format(
                bookmark_id=bookmark[0],
                bookmark_name=bookmark[1],
                anchor_id=anchor[0] if anchor else "",
                anchor_name=anchor[1] if anchor else "",
                text=_escape_xml(unit.value),
            ))
            continue
        paragraphs.append(_paragraph_xml(unit.value, bookmark=bookmark, anchor=anchor))
    _write_docx_xml(path, paragraphs)


def _write_docx_xml(path: Path, paragraphs: list[str]) -> None:
    xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:document xmlns:w="{W}" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"><w:body>'
        + "".join(paragraphs)
        + "</w:body></w:document>"
    )
    with zipfile.ZipFile(path, "w") as package:
        package.writestr("word/document.xml", xml)


def _stage_bookmark(path: Path, prefix: str) -> tuple[str, str]:
    with zipfile.ZipFile(path) as package:
        root = ElementTree.fromstring(package.read("word/document.xml"))
    for node in root.iter(f"{{{W}}}bookmarkStart"):
        name = node.attrib.get(f"{{{W}}}name", "")
        if name.startswith(prefix):
            return node.attrib[f"{{{W}}}id"], name
    raise AssertionError("missing Stage 78 bookmark")


def _append_document_body_xml(path: Path, fragment: str) -> None:
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(path.with_suffix(".rewritten"), "w") as destination:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == "word/document.xml":
                payload = payload.replace(b"</w:body>", fragment.encode("utf-8") + b"</w:body>")
            destination.writestr(info, payload)
    path.with_suffix(".rewritten").replace(path)


def _write_html(
    path: Path,
    units: list[validation.PathwayVisibleUnit],
    *,
    mutate: dict[int, str] | None = None,
    hidden: bool = False,
    wrapper: str = "body",
    wrong_section_index: int | None = None,
    outside_section_index: int | None = None,
    wrong_chapter_index: int | None = None,
    nested_section_index: int | None = None,
    duplicate_section_frame: bool = False,
    missing_section_frame: bool = False,
    unknown_section_frame: bool = False,
    malformed_section: bool = False,
    wrong_container_index: int | None = None,
    duplicate_unit_frame_index: int | None = None,
    unknown_unit_frame_index: int | None = None,
    expected_text_outside: bool = False,
) -> None:
    if wrapper == "comment":
        body = "".join(f"<!-- {html.escape(unit.value)} -->" for unit in units)
        path.write_text(f'<!doctype html><html lang="en"><head><title>ignored</title></head><body>{body}</body></html>', encoding="utf-8")
        return
    if wrapper in {"script", "style", "template", "head", "attribute"}:
        joined = " ".join(html.escape(unit.value, quote=True) for unit in units)
        body = f"<{wrapper}>{joined}</{wrapper}>" if wrapper != "attribute" else f'<div title="{joined}"></div>'
        path.write_text(f'<!doctype html><html lang="en"><head><title>ignored</title></head><body>{body}</body></html>', encoding="utf-8")
        return
    fragments: list[str] = []
    opened_article = False
    opened_section = False
    current_section = ""
    duplicate_frame = units[0].framing_identity if duplicate_unit_frame_index is not None else None
    for index, unit in enumerate(units):
        text = html.escape((mutate or {}).get(index, unit.value))
        style = ' style="display:none"' if hidden else ""
        if unit.field_identity == "chapter":
            chapter_id = "stage78-wrong-chapter" if wrong_chapter_index is not None else unit.framing_identity
            fragments.append(f'<article class="chapter" id="{html.escape(chapter_id, quote=True)}"><h1>{text}</h1>')
            opened_article = True
            continue
        if unit.unit_kind == "heading":
            if opened_section:
                fragments.append("</section>")
                opened_section = False
            section_id = unit.framing_identity
            if missing_section_frame:
                section_id = ""
            if unknown_section_frame:
                section_id = "stage78-section-unknown"
            fragments.append(f'<section class="section" id="{html.escape(section_id, quote=True)}"><h2>{text}</h2>')
            opened_section = True
            current_section = unit.framing_identity
            if duplicate_section_frame:
                fragments.append(f'<section class="section" id="{html.escape(section_id, quote=True)}"><h2>{text}</h2></section>')
            if malformed_section:
                fragments.append("</article>")
            continue
        frame = unit.framing_identity
        if duplicate_unit_frame_index == index and duplicate_frame is not None:
            frame = duplicate_frame
        if unknown_unit_frame_index == index:
            frame = "stage78-unknown-unit"
        if wrong_section_index == index:
            fragments.append("</section>")
            fragments.append('<section class="section" id="stage78-section-wrong"><h2>Wrong section</h2>')
            opened_section = True
            current_section = "stage78-section-wrong"
        if outside_section_index == index and opened_section:
            fragments.append("</section>")
            opened_section = False
        if nested_section_index == index:
            fragments.append('<section class="section" id="stage78-section-nested"><h2>Nested section</h2>')
            current_section = "stage78-section-nested"
        tag = "div" if wrong_container_index == index else "p"
        fragments.append(f'<{tag} class="paragraph-body" id="{html.escape(frame, quote=True)}"{style}>{text}</{tag}>')
        if nested_section_index == index:
            fragments.append("</section>")
            current_section = ""
        if outside_section_index == index and opened_article and current_section:
            fragments.append(f'<section class="section" id="{html.escape(current_section, quote=True)}">')
            opened_section = True
    if expected_text_outside:
        fragments.append(f"<p>{html.escape(units[-1].value)}</p>")
    if opened_section:
        fragments.append("</section>")
    if opened_article:
        fragments.append("</article>")
    body = "".join(fragments)
    path.write_text(f'<!doctype html><html lang="en"><head><title>ignored</title></head><body>{body}</body></html>', encoding="utf-8")


def _pdf_extraction(units: list[validation.PathwayVisibleUnit], *, mutate: dict[int, str] | None = None) -> validation.PdfTextExtraction:
    text = "\n".join((mutate or {}).get(index, unit.value) for index, unit in enumerate(units))
    return validation.PdfTextExtraction(text=validation.normalize_text(text), raw_text=text, status="available", backend="test")


def _pdf_raw_extraction(text: str) -> validation.PdfTextExtraction:
    return validation.PdfTextExtraction(text=validation.normalize_text(text), raw_text=text, status="available", backend="test")


def _unit_with_value(unit: validation.PathwayVisibleUnit, value: str) -> validation.PathwayVisibleUnit:
    return validation.PathwayVisibleUnit(
        record_identity=unit.record_identity,
        unit_kind=unit.unit_kind,
        field_identity=unit.field_identity,
        ordinal=unit.ordinal,
        value=value,
        framing_identity=unit.framing_identity,
        sequence_position=unit.sequence_position,
    )


def _spec_digest(spec: dict) -> str:
    return hashlib.sha256(json.dumps(spec, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")).hexdigest()


def _trusted_pdf_conversion(
    model: dict,
    spec: dict,
    qualification: dict | None,
    docx_path: Path,
    pdf_path: Path,
    *,
    governed_job_id: int = 417,
    governed_attempt_count: int = 1,
    retry_of_job_id: int | None = None,
    converter: str = "test-converter",
) -> dict[str, str]:
    return validation.pathway_pdf_conversion_authority_record(
        model, spec, qualification, governed_job_id=governed_job_id,
        governed_attempt_count=governed_attempt_count, retry_of_job_id=retry_of_job_id,
        docx_sha256=hashlib.sha256(docx_path.read_bytes()).hexdigest(),
        pdf_sha256=hashlib.sha256(pdf_path.read_bytes()).hexdigest(),
        converter_identity="test-converter", converter_version=converter,
        created_at="2026-09-09T12:00:00Z",
    )


class Stage78BPathwayOutputEquivalenceTests(unittest.TestCase):
    def test_canonical_derivation_is_complete_deterministic_and_independent(self) -> None:
        model = _model()
        spec = _spec(model)
        qualification = _qualification()
        units = _units(model, spec, qualification)
        self.assertEqual(units, _units(copy.deepcopy(model), copy.deepcopy(spec), copy.deepcopy(qualification)))
        self.assertEqual(set(KIND_TO_SECTION), set(adapter.PATHWAY_RENDER_OBJECT_KINDS))
        row_values = [unit.value for unit in units if unit.field_identity == "pathway_row"]
        for kind in KIND_TO_SECTION:
            self.assertTrue(any(f"Kind: {kind}" in value for value in row_values), kind)
        self.assertTrue(any(unit.field_identity == "qualification_disclosure" for unit in units))

    def test_invalid_authority_and_qualification_fail_closed(self) -> None:
        model = _model()
        bad_spec = _spec(model)
        bad_spec["pathway_projection"]["projection_digest"] = "c" * 64
        result, audit = validation.validate_pathway_output_equivalence(model, specification=bad_spec)
        self.assertFalse(result.ok)
        self.assertEqual(audit.failures[0]["reason"], "invalid_canonical_authority_input")
        result, audit = validation.validate_pathway_output_equivalence(model, specification=_spec(model), governance_qualification={"disclosure": "too little"})
        self.assertFalse(result.ok)
        self.assertEqual(audit.failures[0]["reason"], "invalid_canonical_authority_input")

    def test_docx_and_html_exact_structural_sequence_pass(self) -> None:
        model = _model()
        spec = _spec(model)
        units = _units(model, spec)
        with tempfile.TemporaryDirectory(prefix="stage78b2b1-") as directory:
            root = Path(directory)
            docx_path = root / "report.docx"
            html_path = root / "report.html"
            _write_framed_docx(docx_path, units)
            _write_html(html_path, units)
            result, audit = validation.validate_pathway_output_equivalence(model, specification=spec, docx_path=docx_path, html_path=html_path)
        self.assertTrue(result.ok, audit.failures)

    def test_renderer_produced_pathway_docx_contains_valid_frames(self) -> None:
        model = _model()
        spec = _spec(model)
        with tempfile.TemporaryDirectory(prefix="stage78b2b1-rendered-") as directory:
            docx_path = Path(directory) / "rendered.docx"
            book = adapter.make_book(spec, pathway_render_model=model)
            adapter.DocxRenderer(adapter.EffectiveTheme(theme=adapter.HANDBOOK_THEME, publication_profile=adapter.PUBLICATION_PROFILES["digital"], page=adapter.HANDBOOK_THEME.page, title_page=adapter.HANDBOOK_THEME.title_page, volume_page=adapter.HANDBOOK_THEME.volume_page, chapter_opening=adapter.HANDBOOK_THEME.chapter_opening)).render(book, docx_path)
            adapter._apply_pathway_docx_heading_unit_frames(docx_path, spec, None, model)
            result, audit = validation.validate_pathway_output_equivalence(model, specification=spec, docx_path=docx_path)
        self.assertTrue(result.ok, audit.failures)

    def test_docx_and_html_missing_extra_duplicate_reordered_or_mutated_units_fail(self) -> None:
        model = _model()
        spec = _spec(model)
        units = _units(model, spec)
        cases = {
            "missing": units[:-1],
            "extra": units + [units[-1]],
            "duplicate": units[:4] + [units[3]] + units[4:],
            "reordered": [units[1], units[0]] + units[2:],
            "punctuation": units,
            "identifier": units,
            "substring": units,
            "blank": units,
        }
        mutations = {
            "punctuation": {3: units[3].value.replace(":", ";", 1)},
            "identifier": {4: units[4].value.replace("CR-STAGE78", "CR-STAGE79")},
            "substring": {5: units[5].value[:20]},
            "blank": {6: " "},
        }
        with tempfile.TemporaryDirectory(prefix="stage78b2b1-bad-") as directory:
            root = Path(directory)
            for name, candidate_units in cases.items():
                with self.subTest(name=name):
                    docx_path = root / f"{name}.docx"
                    html_path = root / f"{name}.html"
                    _write_framed_docx(docx_path, candidate_units, authority_units=units, mutate_text=mutations.get(name))
                    _write_html(html_path, candidate_units, mutate=mutations.get(name))
                    result, audit = validation.validate_pathway_output_equivalence(model, specification=spec, docx_path=docx_path, html_path=html_path)
                    self.assertFalse(result.ok)
                    self.assertTrue(audit.failures)

    def test_cross_row_substitution_and_identical_twin_multiplicity_fail(self) -> None:
        model = _model()
        spec = _spec(model)
        units = _units(model, spec)
        first_row = next(index for index, unit in enumerate(units) if unit.field_identity == "pathway_row")
        second_row = first_row + 1
        with tempfile.TemporaryDirectory(prefix="stage78b2b1-cross-") as directory:
            docx_path = Path(directory) / "cross.docx"
            _write_framed_docx(docx_path, units, mutate_text={first_row: units[second_row].value})
            result, audit = validation.validate_pathway_output_equivalence(model, specification=spec, docx_path=docx_path)
        self.assertFalse(result.ok)
        self.assertTrue(any(item["reason"] in {"value_mismatch", "multiplicity_mismatch"} for item in audit.failures))

    def test_html_non_visible_or_unframed_occurrences_fail(self) -> None:
        model = _model()
        spec = _spec(model)
        units = _units(model, spec)
        with tempfile.TemporaryDirectory(prefix="stage78b2b1-html-") as directory:
            root = Path(directory)
            for wrapper in ("comment", "script", "style", "template", "head", "attribute"):
                with self.subTest(wrapper=wrapper):
                    html_path = root / f"{wrapper}.html"
                    _write_html(html_path, units, wrapper=wrapper)
                    result, audit = validation.validate_pathway_output_equivalence(model, specification=spec, html_path=html_path)
                    self.assertFalse(result.ok)
                    self.assertTrue(audit.failures)
            hidden = root / "hidden.html"
            _write_html(hidden, units, hidden=True)
            result, audit = validation.validate_pathway_output_equivalence(model, specification=spec, html_path=hidden)
            self.assertFalse(result.ok)

    def test_html_structural_ancestry_attacks_fail(self) -> None:
        model = _model()
        spec = _spec(model)
        units = _units(model, spec)
        first_body = next(index for index, unit in enumerate(units) if unit.unit_kind == "body")
        with tempfile.TemporaryDirectory(prefix="stage78b2b1-html-ancestry-") as directory:
            root = Path(directory)
            cases = {
                "wrong_section": {"wrong_section_index": first_body},
                "outside_section": {"outside_section_index": first_body},
                "wrong_chapter": {"wrong_chapter_index": first_body},
                "nested_section": {"nested_section_index": first_body},
                "duplicate_section": {"duplicate_section_frame": True},
                "missing_section": {"missing_section_frame": True},
                "unknown_section": {"unknown_section_frame": True},
                "malformed_section": {"malformed_section": True},
                "wrong_container": {"wrong_container_index": first_body},
                "duplicate_unit": {"duplicate_unit_frame_index": first_body},
                "unknown_unit": {"unknown_unit_frame_index": first_body},
                "expected_text_outside": {"expected_text_outside": True},
            }
            for name, options in cases.items():
                with self.subTest(name=name):
                    html_path = root / f"{name}.html"
                    _write_html(html_path, units, **options)
                    result, audit = validation.validate_pathway_output_equivalence(model, specification=spec, html_path=html_path)
                    self.assertFalse(result.ok)
                    self.assertTrue(audit.failures)

    def test_docx_non_body_hidden_and_field_occurrences_fail(self) -> None:
        model = _model()
        spec = _spec(model)
        units = _units(model, spec)
        with tempfile.TemporaryDirectory(prefix="stage78b2b1-docx-") as directory:
            root = Path(directory)
            property_only = root / "property.docx"
            document = Document()
            document.core_properties.comments = units[0].value
            document.save(property_only)
            result, audit = validation.validate_pathway_output_equivalence(model, specification=spec, docx_path=property_only)
            self.assertFalse(result.ok)
            hidden = root / "hidden.docx"
            document = Document()
            for unit in units:
                document.add_paragraph().add_run(unit.value).font.hidden = True
            document.save(hidden)
            result, audit = validation.validate_pathway_output_equivalence(model, specification=spec, docx_path=hidden)
            self.assertFalse(result.ok)
            source = root / "source.docx"
            field = root / "field.docx"
            _write_framed_docx(source, units[:1])
            with zipfile.ZipFile(source) as src, zipfile.ZipFile(field, "w") as dst:
                for info in src.infolist():
                    data = src.read(info.filename)
                    if info.filename == "word/document.xml":
                        data = data.replace(b"<w:t>Chapter 1", b"<w:instrText>Chapter 1")
                        data = data.replace(b"</w:t>", b"</w:instrText>", 1)
                    dst.writestr(info, data)
            result, audit = validation.validate_pathway_output_equivalence(model, specification=spec, docx_path=field)
            self.assertFalse(result.ok)

    def test_docx_structural_frame_attacks_fail(self) -> None:
        model = _model()
        spec = _spec(model)
        units = _units(model, spec)
        first_body = next(index for index, unit in enumerate(units) if unit.unit_kind == "body")
        second_body = first_body + 1
        first_pathway_row = next(index for index, unit in enumerate(units) if unit.field_identity == "pathway_row")
        with tempfile.TemporaryDirectory(prefix="stage78b2b1-docx-frames-") as directory:
            root = Path(directory)
            cases = {
                "unframed": lambda path: _write_docx(path, units),
                "wrong_frame": lambda path: _write_framed_docx(path, units, mutate_token={first_body: validation.pathway_docx_frame_contract(units)["unit_token_by_frame"][units[second_body].framing_identity]}),
                "wrong_section": lambda path: _write_framed_docx(path, [*units[:2], units[first_pathway_row], *units[2:first_pathway_row], *units[first_pathway_row + 1:]]),
                "table_cell": lambda path: _write_framed_docx(path, units, table_index=first_body),
                "duplicate_name": lambda path: _write_framed_docx(path, units, duplicate_name_index=first_body),
                "duplicate_id": lambda path: _write_framed_docx(path, units, duplicate_id_index=first_body),
                "spanning": lambda path: _write_framed_docx(path, units, spanning_pair=(first_body, second_body)),
                "overlapping": lambda path: _write_framed_docx(path, units, overlapping_pair=(first_body, second_body)),
                "missing_end": lambda path: _write_framed_docx(path, units, missing_end_index=first_body),
                "hidden_only": lambda path: _write_framed_docx(path, units, hidden_index=first_body),
                "deleted_only": lambda path: _write_framed_docx(path, units, deleted_index=first_body),
                "instruction_only": lambda path: _write_framed_docx(path, units, instruction_index=first_body),
                "metadata_only": lambda path: _write_framed_docx(path, units, metadata_only_index=first_body),
                "unknown_extra": lambda path: _write_framed_docx(path, units, unknown_extra=True),
            }
            for name, writer in cases.items():
                with self.subTest(name=name):
                    docx_path = root / f"{name}.docx"
                    writer(docx_path)
                    result, audit = validation.validate_pathway_output_equivalence(model, specification=spec, docx_path=docx_path)
                    self.assertFalse(result.ok)
                    self.assertTrue(audit.failures)

    def test_docx_document_global_bookmark_graph_rejects_collisions_and_malformed_ranges(self) -> None:
        model = _model()
        spec = _spec(model)
        units = _units(model, spec)
        with tempfile.TemporaryDirectory(prefix="stage78b2b1-docx-global-graph-") as directory:
            root = Path(directory)
            baseline = root / "baseline.docx"
            _write_framed_docx(baseline, units)
            visible_id, _ = _stage_bookmark(baseline, validation.PATHWAY_DOCX_VISIBLE_PREFIX)
            structural_id, _ = _stage_bookmark(baseline, validation.PATHWAY_DOCX_STRUCTURAL_PREFIX)
            cases = {
                "visible_id_collision": f'<w:p><w:bookmarkStart w:id="{visible_id}" w:name="other-visible"/><w:bookmarkEnd w:id="{visible_id}"/></w:p>',
                "structural_id_collision": f'<w:p><w:bookmarkStart w:id="{structural_id}" w:name="other-structural"/><w:bookmarkEnd w:id="{structural_id}"/></w:p>',
                "duplicate_non_stage_start": '<w:p><w:bookmarkStart w:id="900" w:name="other-a"/><w:bookmarkStart w:id="900" w:name="other-b"/><w:bookmarkEnd w:id="900"/></w:p>',
                "duplicate_non_stage_end": '<w:p><w:bookmarkStart w:id="901" w:name="other"/><w:bookmarkEnd w:id="901"/><w:bookmarkEnd w:id="901"/></w:p>',
                "mixed_stage_non_stage_range": f'<w:p><w:bookmarkStart w:id="{visible_id}" w:name="other-mixed"/><w:bookmarkEnd w:id="{visible_id}"/></w:p>',
                "start_without_end": '<w:p><w:bookmarkStart w:id="902" w:name="other"/></w:p>',
                "end_without_start": '<w:p><w:bookmarkEnd w:id="903"/></w:p>',
                "ambiguous_crossed_pairing": '<w:p><w:bookmarkStart w:id="904" w:name="other-a"/><w:bookmarkStart w:id="905" w:name="other-b"/><w:bookmarkEnd w:id="904"/><w:bookmarkEnd w:id="905"/></w:p>',
                "elsewhere_collision": f'<w:p><w:bookmarkStart w:id="{visible_id}" w:name="outside-stage78"/><w:bookmarkEnd w:id="{visible_id}"/></w:p>',
                "malformed_global_graph_with_correct_stage78": '<w:p><w:bookmarkStart w:id="906" w:name="other-a"/><w:bookmarkStart w:id="906" w:name="other-b"/><w:bookmarkEnd w:id="906"/></w:p>',
            }
            for name, fragment in cases.items():
                with self.subTest(name=name):
                    docx_path = root / f"{name}.docx"
                    docx_path.write_bytes(baseline.read_bytes())
                    _append_document_body_xml(docx_path, fragment)
                    result, audit = validation.validate_pathway_output_equivalence(model, specification=spec, docx_path=docx_path)
                    self.assertFalse(result.ok)
                    self.assertEqual(audit.failures, [{"format": "docx", "field": "artifact_structure", "reason": "malformed_artifact_structure"}])
                    self.assertNotIn(units[0].value, str(audit.failures))
            valid_non_stage = root / "valid-non-stage.docx"
            valid_non_stage.write_bytes(baseline.read_bytes())
            _append_document_body_xml(valid_non_stage, '<w:p><w:bookmarkStart w:id="907" w:name="other-valid"/><w:r><w:t>Independent bookmark</w:t></w:r><w:bookmarkEnd w:id="907"/></w:p>')
            result, audit = validation.validate_pathway_output_equivalence(model, specification=spec, docx_path=valid_non_stage)
        self.assertTrue(result.ok, audit.failures)

    def test_docx_ancestor_visibility_and_bookmark_span_attacks_fail(self) -> None:
        model = _model()
        spec = _spec(model)
        units = _units(model, spec)
        first_body = next(index for index, unit in enumerate(units) if unit.unit_kind == "body")
        excluded_cases = {
            "move_from": '<w:p><w:bookmarkStart w:id="{bookmark_id}" w:name="{bookmark_name}"/><w:moveFrom><w:r><w:t>{text}</w:t></w:r></w:moveFrom><w:bookmarkEnd w:id="{bookmark_id}"/></w:p>',
            "deleted_container": '<w:p><w:bookmarkStart w:id="{bookmark_id}" w:name="{bookmark_name}"/><w:del><w:r><w:t>{text}</w:t></w:r></w:del><w:bookmarkEnd w:id="{bookmark_id}"/></w:p>',
            "deleted_text": '<w:p><w:bookmarkStart w:id="{bookmark_id}" w:name="{bookmark_name}"/><w:r><w:delText>{text}</w:delText></w:r><w:bookmarkEnd w:id="{bookmark_id}"/></w:p>',
            "instruction_text": '<w:p><w:bookmarkStart w:id="{bookmark_id}" w:name="{bookmark_name}"/><w:r><w:instrText>{text}</w:instrText></w:r><w:bookmarkEnd w:id="{bookmark_id}"/></w:p>',
            "vanished_run": '<w:p><w:bookmarkStart w:id="{bookmark_id}" w:name="{bookmark_name}"/><w:r><w:rPr><w:vanish/></w:rPr><w:t>{text}</w:t></w:r><w:bookmarkEnd w:id="{bookmark_id}"/></w:p>',
            "web_hidden_run": '<w:p><w:bookmarkStart w:id="{bookmark_id}" w:name="{bookmark_name}"/><w:r><w:rPr><w:webHidden/></w:rPr><w:t>{text}</w:t></w:r><w:bookmarkEnd w:id="{bookmark_id}"/></w:p>',
            "alternate_content": '<w:p><w:bookmarkStart w:id="{bookmark_id}" w:name="{bookmark_name}"/><mc:AlternateContent><mc:Choice><w:r><w:t>{text}</w:t></w:r></mc:Choice><mc:Fallback><w:r><w:t>{text}</w:t></w:r></mc:Fallback></mc:AlternateContent><w:bookmarkEnd w:id="{bookmark_id}"/></w:p>',
            "text_box": '<w:p><w:bookmarkStart w:id="{bookmark_id}" w:name="{bookmark_name}"/><w:drawing><w:txbxContent><w:r><w:t>{text}</w:t></w:r></w:txbxContent></w:drawing><w:bookmarkEnd w:id="{bookmark_id}"/></w:p>',
            "unsupported_wrapper": '<w:p><w:bookmarkStart w:id="{bookmark_id}" w:name="{bookmark_name}"/><w:unsupportedTextWrapper><w:r><w:t>{text}</w:t></w:r></w:unsupportedTextWrapper><w:bookmarkEnd w:id="{bookmark_id}"/></w:p>',
            "outside_bookmark_visible_copy": '<w:p><w:bookmarkStart w:id="{bookmark_id}" w:name="{bookmark_name}"/><w:moveFrom><w:r><w:t>{text}</w:t></w:r></w:moveFrom><w:bookmarkEnd w:id="{bookmark_id}"/><w:r><w:t>{text}</w:t></w:r></w:p>',
            "nested_deleted_move_from": '<w:p><w:bookmarkStart w:id="{bookmark_id}" w:name="{bookmark_name}"/><w:del><w:moveFrom><w:r><w:t>{text}</w:t></w:r></w:moveFrom></w:del><w:bookmarkEnd w:id="{bookmark_id}"/></w:p>',
            "hidden_under_hyperlink": '<w:p><w:bookmarkStart w:id="{bookmark_id}" w:name="{bookmark_name}"/><w:hyperlink><w:r><w:rPr><w:vanish/></w:rPr><w:t>{text}</w:t></w:r></w:hyperlink><w:bookmarkEnd w:id="{bookmark_id}"/></w:p>',
            "hidden_under_content_control": '<w:p><w:bookmarkStart w:id="{bookmark_id}" w:name="{bookmark_name}"/><w:sdt><w:sdtContent><w:r><w:rPr><w:vanish/></w:rPr><w:t>{text}</w:t></w:r></w:sdtContent></w:sdt><w:bookmarkEnd w:id="{bookmark_id}"/></w:p>',
        }
        value = units[first_body].value
        midpoint = len(value) // 2
        excluded_cases.update({
            "visible_prefix_move_from_suffix": f'<w:p><w:bookmarkStart w:id="{{bookmark_id}}" w:name="{{bookmark_name}}"/><w:r><w:t>{_escape_xml(value[:midpoint])}</w:t></w:r><w:moveFrom><w:r><w:t>{_escape_xml(value[midpoint:])}</w:t></w:r></w:moveFrom><w:bookmarkEnd w:id="{{bookmark_id}}"/></w:p>',
            "move_from_prefix_visible_suffix": f'<w:p><w:bookmarkStart w:id="{{bookmark_id}}" w:name="{{bookmark_name}}"/><w:moveFrom><w:r><w:t>{_escape_xml(value[:midpoint])}</w:t></w:r></w:moveFrom><w:r><w:t>{_escape_xml(value[midpoint:])}</w:t></w:r><w:bookmarkEnd w:id="{{bookmark_id}}"/></w:p>',
            "excluded_between_visible_runs": '<w:p><w:bookmarkStart w:id="{bookmark_id}" w:name="{bookmark_name}"/><w:r><w:t>Alpha</w:t></w:r><w:moveFrom><w:r><w:t> hidden </w:t></w:r></w:moveFrom><w:r><w:t>Beta</w:t></w:r><w:bookmarkEnd w:id="{bookmark_id}"/></w:p>',
        })
        with tempfile.TemporaryDirectory(prefix="stage78b2b1-docx-ancestors-") as directory:
            root = Path(directory)
            for name, paragraph_xml in excluded_cases.items():
                with self.subTest(name=name):
                    docx_path = root / f"{name}.docx"
                    case_units = units
                    if name == "excluded_between_visible_runs":
                        case_units = [*units]
                        case_units[first_body] = _unit_with_value(units[first_body], "Alpha hidden Beta")
                    _write_framed_docx_with_override(docx_path, case_units, first_body, paragraph_xml)
                    result, audit = validation.validate_pathway_output_equivalence(model, specification=spec, docx_path=docx_path)
                    self.assertFalse(result.ok)
                    self.assertTrue(audit.failures)

    def test_docx_supported_visible_ancestors_and_controls_pass(self) -> None:
        model = _model()
        spec = _spec(model)
        units = _units(model, spec)
        first_body = next(index for index, unit in enumerate(units) if unit.unit_kind == "body")
        words = units[first_body].value.split(" ")
        chunks = [
            " ".join(words[:3]),
            " " + " ".join(words[3:6]),
            " ".join(words[6:9]),
            " ".join(words[9:12]),
            " " + " ".join(words[12:15]),
            " ".join(words[15:]),
        ]
        paragraph_xml = (
            '<w:p><w:bookmarkStart w:id="{bookmark_id}" w:name="{bookmark_name}"/>'
            f'<w:r><w:t>{_escape_xml(chunks[0])}</w:t></w:r>'
            f'<w:hyperlink><w:r><w:t>{_escape_xml(chunks[1])}</w:t></w:r></w:hyperlink>'
            f'<w:ins><w:r><w:tab/><w:t>{_escape_xml(chunks[2])}</w:t></w:r></w:ins>'
            f'<w:moveTo><w:r><w:br/><w:t>{_escape_xml(chunks[3])}</w:t></w:r></w:moveTo>'
            f'<w:sdt><w:sdtContent><w:r><w:t>{_escape_xml(chunks[4])}</w:t></w:r></w:sdtContent></w:sdt>'
            f'<w:r><w:cr/><w:t>{_escape_xml(chunks[5])}</w:t></w:r>'
            '<w:bookmarkEnd w:id="{bookmark_id}"/></w:p>'
        )
        with tempfile.TemporaryDirectory(prefix="stage78b2b1-docx-visible-") as directory:
            docx_path = Path(directory) / "visible.docx"
            _write_framed_docx_with_override(docx_path, units, first_body, paragraph_xml)
            result, audit = validation.validate_pathway_output_equivalence(model, specification=spec, docx_path=docx_path)
        self.assertTrue(result.ok, audit.failures)

    def test_docx_duplicate_visible_values_require_distinct_frames(self) -> None:
        model = _model()
        first = model["sections"][1]["rows"][0]
        second = model["sections"][1]["rows"][1]
        second["object_kind"] = first["object_kind"]
        second["governed_logical_identity"] = first["governed_logical_identity"]
        second["category"] = first["category"]
        second["status"] = first["status"]
        second["represented_time"] = first["represented_time"]
        second["recorded_at"] = first["recorded_at"]
        second["chronology"] = copy.deepcopy(first["chronology"])
        second["ownership_path"] = first["ownership_path"]
        second["source_authority_key"] = first["source_authority_key"]
        second["row_authority_digest"] = first["row_authority_digest"]
        second["parent_governed_identity"] = first["parent_governed_identity"]
        second["endpoint_identities"] = copy.deepcopy(first["endpoint_identities"])
        second["contestation"] = copy.deepcopy(first["contestation"])
        second["supersession"] = copy.deepcopy(first["supersession"])
        second["reliance"] = copy.deepcopy(first["reliance"])
        second["limitations"] = copy.deepcopy(first["limitations"])
        second["does_not_establish"] = copy.deepcopy(first["does_not_establish"])
        spec = _spec(model)
        units = _units(model, spec)
        with tempfile.TemporaryDirectory(prefix="stage78b2b1-docx-twins-") as directory:
            docx_path = Path(directory) / "twins.docx"
            _write_framed_docx(docx_path, units)
            result, audit = validation.validate_pathway_output_equivalence(model, specification=spec, docx_path=docx_path)
        self.assertTrue(result.ok, audit.failures)

    def test_pdf_trusted_conversion_binding_and_exact_sequence_pass(self) -> None:
        model = _model()
        spec = _spec(model)
        qualification = _qualification()
        units = _units(model, spec, qualification)
        with tempfile.TemporaryDirectory(prefix="stage78b2b1-pdf-") as directory:
            root = Path(directory)
            docx_path = root / "report.docx"
            pdf_path = root / "report.pdf"
            _write_framed_docx(docx_path, units)
            pdf_path.write_bytes(b"%PDF-1.7\nsynthetic\n%%EOF\n")
            conversion = _trusted_pdf_conversion(model, spec, qualification, docx_path, pdf_path)
            with patch.object(validation, "extract_pdf_text_result", return_value=_pdf_extraction(units)):
                result, audit = validation.validate_pathway_output_equivalence(model, specification=spec, governance_qualification=qualification, docx_path=docx_path, pdf_path=pdf_path, trusted_pdf_conversion=conversion, governed_job_id=417, governed_attempt_count=1)
        self.assertTrue(result.ok, audit.failures)
        self.assertIn("trusted pinned DOCX conversion", audit.pdf_claim)
        self.assertIn("not independently proven", audit.pdf_claim)

    def test_pdf_conversion_authority_digest_is_deterministic_and_attempt_bound(self) -> None:
        model = _model(); spec = _spec(model); qualification = _qualification()
        with tempfile.TemporaryDirectory(prefix="stage78b2b1-pdf-authority-") as directory:
            root = Path(directory); docx_path = root / "report.docx"; pdf_path = root / "report.pdf"
            _write_framed_docx(docx_path, _units(model, spec, qualification)); pdf_path.write_bytes(b"%PDF-1.7\nsynthetic\n%%EOF\n")
            first = _trusted_pdf_conversion(model, spec, qualification, docx_path, pdf_path, governed_attempt_count=1)
            second = _trusted_pdf_conversion(model, spec, qualification, docx_path, pdf_path, governed_attempt_count=1)
            retry = _trusted_pdf_conversion(model, spec, qualification, docx_path, pdf_path, governed_attempt_count=2, retry_of_job_id=416)
        self.assertEqual(first, second)
        self.assertNotEqual(first["conversion_record_digest"], retry["conversion_record_digest"])
        self.assertEqual(retry["retry_predecessor_identity"], "stage77-report-job:416")
        malformed = dict(first); malformed["attempt_identity"] = "stage77-report-job:417:attempt:2"
        self.assertNotEqual(malformed["conversion_record_digest"], validation._pathway_sha256({key: value for key, value in malformed.items() if key != "conversion_record_digest"}))

    def test_pdf_binding_backend_and_extraction_fail_closed(self) -> None:
        model = _model()
        spec = _spec(model)
        qualification = _qualification()
        units = _units(model, spec, qualification)
        with tempfile.TemporaryDirectory(prefix="stage78b2b1-pdf-bad-") as directory:
            root = Path(directory)
            docx_path = root / "report.docx"
            pdf_path = root / "report.pdf"
            _write_framed_docx(docx_path, units)
            pdf_path.write_bytes(b"%PDF-1.7\nsynthetic\n%%EOF\n")
            good = _trusted_pdf_conversion(model, spec, qualification, docx_path, pdf_path)
            for mutation in ({"source_docx_sha256": "0" * 64}, {"pdf_sha256": "1" * 64}, {"rendering_profile": "wrong"}, {"converter": ""}):
                with self.subTest(mutation=mutation):
                    with patch.object(validation, "extract_pdf_text_result", return_value=_pdf_extraction(units)):
                        result, audit = validation.validate_pathway_output_equivalence(model, specification=spec, governance_qualification=qualification, docx_path=docx_path, pdf_path=pdf_path, trusted_pdf_conversion=dict(good, **mutation), governed_job_id=417, governed_attempt_count=1)
                    self.assertFalse(result.ok)
                    self.assertTrue(any(item["reason"] == "trusted_conversion_binding_failure" for item in audit.failures))
            with patch.object(validation, "extract_pdf_text_result", return_value=validation.PdfTextExtraction(status="unavailable", reason="missing backend")):
                result, audit = validation.validate_pathway_output_equivalence(model, specification=spec, governance_qualification=qualification, docx_path=docx_path, pdf_path=pdf_path, trusted_pdf_conversion=good, governed_job_id=417, governed_attempt_count=1)
            self.assertFalse(result.ok)
            self.assertTrue(any(item["reason"] == "pdf_extraction_backend_unavailable" for item in audit.failures))

    def test_pdf_governance_chain_binding_failures_override_matching_text_and_digests(self) -> None:
        model = _model()
        spec = _spec(model)
        qualification = _qualification()
        units = _units(model, spec, qualification)
        with tempfile.TemporaryDirectory(prefix="stage78b2b1-pdf-chain-") as directory:
            root = Path(directory)
            docx_path = root / "report.docx"
            pdf_path = root / "report.pdf"
            _write_framed_docx(docx_path, units)
            pdf_path.write_bytes(b"%PDF-1.7\nsynthetic\n%%EOF\n")
            good = _trusted_pdf_conversion(model, spec, qualification, docx_path, pdf_path)
            other_model = copy.deepcopy(model)
            other_model["projection_digest"] = "c" * 64
            other_model["sections"][1]["rows"][0]["row_authority_digest"] = "d" * 64
            other_qualification = dict(qualification, disclosure="Different governed qualification.")
            cases = {
                "missing_job_identity": {key: value for key, value in good.items() if key != "job_identity"},
                "wrong_job_identity": dict(good, job_identity="stage77-report-job:418"),
                "malformed_job_identity": dict(good, job_identity="stage77-report-job:0417"),
                "replayed_under_another_job": good,
                "missing_specification_identity": {key: value for key, value in good.items() if key != "specification_identity"},
                "wrong_specification_identity": dict(good, specification_identity="wrong.specification.v1"),
                "missing_specification_digest": {key: value for key, value in good.items() if key != "specification_sha256"},
                "wrong_specification_digest": dict(good, specification_sha256="1" * 64),
                "missing_model_identity": {key: value for key, value in good.items() if key != "render_model_identity"},
                "wrong_model_identity": dict(good, render_model_identity="stage78.other_model.v1"),
                "missing_model_digest": {key: value for key, value in good.items() if key != "render_model_sha256"},
                "different_model_digest": dict(good, render_model_sha256=validation.pathway_trusted_pdf_conversion_binding(other_model, _spec(other_model), qualification, governed_job_id=417, governed_attempt_count=1)["render_model_sha256"]),
                "model_mutated_after_record": good,
                "missing_qualification_identity": {key: value for key, value in good.items() if key != "governance_qualification_identity"},
                "missing_qualification_digest": {key: value for key, value in good.items() if key != "governance_qualification_sha256"},
                "different_qualification_digest": dict(good, governance_qualification_sha256=validation.pathway_trusted_pdf_conversion_binding(model, spec, other_qualification, governed_job_id=417, governed_attempt_count=1)["governance_qualification_sha256"]),
                "qualification_mutated_after_record": good,
                "model_qualification_swapped": dict(good, render_model_sha256=good["governance_qualification_sha256"], governance_qualification_sha256=good["render_model_sha256"]),
                "malformed_digest": dict(good, render_model_sha256=True),
                "extra_field": dict(good, unexpected="extra"),
            }
            for name, conversion in cases.items():
                with self.subTest(name=name):
                    call_model = copy.deepcopy(model)
                    call_spec = copy.deepcopy(spec)
                    call_qualification = copy.deepcopy(qualification)
                    if name == "model_mutated_after_record":
                        call_model["sections"][1]["rows"][0]["status"] = "altered"
                    elif name == "qualification_mutated_after_record":
                        call_qualification["disclosure"] = "Different governed qualification."
                    with patch.object(validation, "extract_pdf_text_result", return_value=_pdf_extraction(units)):
                        expected_job_id = 418 if name == "replayed_under_another_job" else 417
                        result, audit = validation.validate_pathway_output_equivalence(call_model, specification=call_spec, governance_qualification=call_qualification, docx_path=docx_path, pdf_path=pdf_path, trusted_pdf_conversion=conversion, governed_job_id=expected_job_id, governed_attempt_count=1)
                    self.assertFalse(result.ok)
                    self.assertTrue(any(item["reason"] == "trusted_conversion_binding_failure" for item in audit.failures), audit.failures)

    def test_pdf_requires_independent_positive_governed_job_id(self) -> None:
        model = _model()
        spec = _spec(model)
        qualification = _qualification()
        units = _units(model, spec, qualification)
        with tempfile.TemporaryDirectory(prefix="stage78b2b1-pdf-job-id-") as directory:
            root = Path(directory)
            docx_path = root / "report.docx"
            pdf_path = root / "report.pdf"
            _write_framed_docx(docx_path, units)
            pdf_path.write_bytes(b"%PDF-1.7\nsynthetic\n%%EOF\n")
            conversion = _trusted_pdf_conversion(model, spec, qualification, docx_path, pdf_path, governed_job_id=417)
            for governed_job_id in (None, True, 0, -1, "417", 417.0):
                with self.subTest(governed_job_id=governed_job_id):
                    with patch.object(validation, "extract_pdf_text_result", return_value=_pdf_extraction(units)):
                        result, audit = validation.validate_pathway_output_equivalence(
                            model,
                            specification=spec,
                            governance_qualification=qualification,
                            docx_path=docx_path,
                            pdf_path=pdf_path,
                            trusted_pdf_conversion=conversion,
                            governed_job_id=governed_job_id, governed_attempt_count=1,
                        )
                    self.assertFalse(result.ok)
                    self.assertTrue(any(item["reason"] == "trusted_conversion_binding_failure" for item in audit.failures), audit.failures)

    def test_pdf_missing_extra_duplicate_reordered_and_metadata_only_fail(self) -> None:
        model = _model()
        spec = _spec(model)
        qualification = _qualification()
        units = _units(model, spec, qualification)
        with tempfile.TemporaryDirectory(prefix="stage78b2b1-pdf-units-") as directory:
            root = Path(directory)
            docx_path = root / "report.docx"
            pdf_path = root / "report.pdf"
            _write_framed_docx(docx_path, units)
            pdf_path.write_bytes(b"%PDF-1.7\nsynthetic\n%%EOF\n")
            conversion = _trusted_pdf_conversion(model, spec, qualification, docx_path, pdf_path)
            for name, extracted in {"missing": units[:-1], "extra": units + [units[-1]], "duplicate": units[:3] + [units[2]] + units[3:], "reordered": [units[1], units[0]] + units[2:], "metadata_only": []}.items():
                with self.subTest(name=name):
                    with patch.object(validation, "extract_pdf_text_result", return_value=_pdf_extraction(extracted)):
                        result, audit = validation.validate_pathway_output_equivalence(model, specification=spec, governance_qualification=qualification, docx_path=docx_path, pdf_path=pdf_path, trusted_pdf_conversion=conversion, governed_job_id=417, governed_attempt_count=1)
                    self.assertFalse(result.ok)
                    self.assertTrue(audit.failures)

    def test_pdf_rejects_precanonical_content_before_first_unit(self) -> None:
        model = _model()
        spec = _spec(model)
        qualification = _qualification()
        units = _units(model, spec, qualification)
        canonical = "\n".join(unit.value for unit in units)
        running_header = "Structured \u00b7 Traceable \u00b7 Governed \u00b7 7 EVIDENCE-LED GOVERNANCE"
        cases = {
            "arbitrary_text": f"unauthorized preface\n{canonical}",
            "governed_looking_row": f"{units[3].value}\n{canonical}",
            "duplicate_later_unit": f"{units[-1].value}\n{canonical}",
            "other_record_unit": f"{units[6].value}\n{canonical}",
            "near_match_header": f"Structured \u00b7 Traceable \u00b7 Governed \u00b7 X EVIDENCE-LED GOVERNANCE\n{canonical}",
            "malformed_running_header": f"Structured \u00b7 Traceable \u00b7 Governed \u00b7 7 Evidence-Led Governance\n{canonical}",
            "partial_running_header": f"Structured \u00b7 Traceable \u00b7 Governed\n{canonical}",
            "header_like_semantic_content": f"Semantic note: {running_header} is governed content\n{canonical}",
            "whitespace_then_injected": f"  \n\t\nunauthorized preface\n{canonical}",
            "form_feed_then_injected": f"\funauthorized preface\n{canonical}",
            "injected_before_authorized_header": f"unauthorized preface\n{running_header}\n{canonical}",
            "same_line_before_first_unit": f"unauthorized preface {units[0].value}\n" + "\n".join(unit.value for unit in units[1:]),
        }
        with tempfile.TemporaryDirectory(prefix="stage78b2b1-pdf-precanonical-") as directory:
            root = Path(directory)
            docx_path = root / "report.docx"
            pdf_path = root / "report.pdf"
            _write_framed_docx(docx_path, units)
            pdf_path.write_bytes(b"%PDF-1.7\nsynthetic\n%%EOF\n")
            conversion = _trusted_pdf_conversion(model, spec, qualification, docx_path, pdf_path)
            for name, raw_text in cases.items():
                with self.subTest(name=name):
                    with patch.object(validation, "extract_pdf_text_result", return_value=_pdf_raw_extraction(raw_text)):
                        with patch.object(validation, "_units_from_triplets", side_effect=AssertionError("triplet synthesis reached after precanonical violation")):
                            result, audit = validation.validate_pathway_output_equivalence(model, specification=spec, governance_qualification=qualification, docx_path=docx_path, pdf_path=pdf_path, trusted_pdf_conversion=conversion, governed_job_id=417, governed_attempt_count=1)
                    self.assertFalse(result.ok)
                    self.assertTrue(any(item["reason"] == "precanonical_content" for item in audit.failures), audit.failures)
                    self.assertNotIn("unauthorized preface", str(audit.failures))
                    self.assertNotIn("Semantic note", str(audit.failures))

            with patch.object(validation, "extract_pdf_text_result", return_value=_pdf_raw_extraction(f"{running_header}\n\n\f{canonical}")):
                result, audit = validation.validate_pathway_output_equivalence(model, specification=spec, governance_qualification=qualification, docx_path=docx_path, pdf_path=pdf_path, trusted_pdf_conversion=conversion, governed_job_id=417, governed_attempt_count=1)
            self.assertTrue(result.ok, audit.failures)

        header_value = f"{running_header} retained as governed text"
        expected = [
            validation.PathwayVisibleUnit(
                record_identity=units[0].record_identity,
                unit_kind=units[0].unit_kind,
                field_identity=units[0].field_identity,
                ordinal=units[0].ordinal,
                value=header_value,
                framing_identity=units[0].framing_identity,
                sequence_position=units[0].sequence_position,
            ),
            *units[1:],
        ]
        pdf_units, failures = validation._pdf_body_units(_pdf_raw_extraction("\n".join(unit.value for unit in expected)), expected)
        self.assertFalse(failures)
        self.assertEqual(pdf_units[0].value, header_value)

    def test_pdf_logical_stream_reconstructs_wrapped_units(self) -> None:
        model = _model()
        spec = _spec(model)
        qualification = _qualification()
        units = _units(model, spec, qualification)
        digest = "b" * 64
        header = "Structured \u00b7 Traceable \u00b7 Governed \u00b7 8 EVIDENCE-LED GOVERNANCE"
        expected = [
            units[0],
            _unit_with_value(units[1], "A long prose value wraps across physical PDF extraction lines without changing semantics."),
            _unit_with_value(units[2], f"Row authority digest: {digest}"),
            _unit_with_value(units[3], "Source URL: https://example.test/pathway/report/alpha-beta-gamma"),
            _unit_with_value(units[4], 'Coverage JSON: {"alpha":["one","two"],"state":"available"}'),
            _unit_with_value(units[5], "The complaint-investigation record remains traceable."),
            _unit_with_value(units[6], "The complaint-investigation record remains traceable."),
        ]
        raw = "\n".join((
            expected[0].value,
            "A long prose value wraps across physical",
            "PDF extraction lines without changing semantics.",
            f"Row authority digest: {digest[:32]}",
            header,
            digest[32:],
            "Source URL: https://example.test/pathway/",
            "report/alpha-beta-gamma",
            'Coverage JSON: {"alpha":["one",',
            '"two"],"state":"available"}',
            "The complaint-",
            "investigation record remains traceable.",
            "\f",
            "The complaint-investigation",
            "record remains traceable.",
        ))
        pdf_units, failures = validation._pdf_body_units(_pdf_raw_extraction(raw), expected)
        self.assertFalse(failures)
        self.assertEqual([unit.value for unit in pdf_units], [unit.value for unit in expected])

    def test_pdf_logical_stream_rejects_unaccounted_prefix_gap_suffix_and_overlap(self) -> None:
        model = _model()
        spec = _spec(model)
        qualification = _qualification()
        units = _units(model, spec, qualification)
        canonical = "\n".join(unit.value for unit in units)
        header_near_match = "Structured \u00b7 Traceable \u00b7 Governed \u00b7 eight EVIDENCE-LED GOVERNANCE"
        cases = {
            "arbitrary_prefix": (f"prefix\n{canonical}", units),
            "same_line_prefix": (f"prefix {units[0].value}\n" + "\n".join(unit.value for unit in units[1:]), units),
            "arbitrary_suffix": (f"{canonical}\nsuffix", units),
            "same_line_suffix": ("\n".join(unit.value for unit in units[:-1]) + f"\n{units[-1].value} suffix", units),
            "extra_page_after_final": (f"{canonical}\n\f\nsuffix", units),
            "arbitrary_line_between": (f"{units[0].value}\ninjected\n" + "\n".join(unit.value for unit in units[1:]), units),
            "same_line_inserted_word": (f"{units[0].value}\n{units[1].value} injected\n" + "\n".join(unit.value for unit in units[2:]), units),
            "punctuation_only_injection": (f"{units[0].value}\n!\n" + "\n".join(unit.value for unit in units[1:]), units),
            "another_record_value": (f"{units[0].value}\n{units[6].value}\n" + "\n".join(unit.value for unit in units[1:]), units),
            "duplicate_earlier_unit": (f"{units[0].value}\n{units[0].value}\n" + "\n".join(unit.value for unit in units[1:]), units),
            "duplicate_later_unit": (f"{units[0].value}\n{units[-1].value}\n" + "\n".join(unit.value for unit in units[1:]), units),
            "running_header_near_match_between": (f"{units[0].value}\n{header_near_match}\n" + "\n".join(unit.value for unit in units[1:]), units),
            "form_feed_adjacent_injection": (f"{units[0].value}\n\f injected\n" + "\n".join(unit.value for unit in units[1:]), units),
            "one_identical_occurrence_for_two_expected": ("duplicate value", [_unit_with_value(units[0], "duplicate value"), _unit_with_value(units[1], "duplicate value")]),
            "duplicate_when_expected_once": ("duplicate value\nduplicate value", [_unit_with_value(units[0], "duplicate value")]),
            "shorter_inside_longer": ("response deadline\ndeadline", [_unit_with_value(units[0], "deadline"), _unit_with_value(units[1], "response deadline")]),
            "longer_then_missing_shorter": ("response deadline", [_unit_with_value(units[0], "response deadline"), _unit_with_value(units[1], "deadline")]),
            "digest_prefix": ("Row authority digest: " + "c" * 32, [_unit_with_value(units[0], "Row authority digest: " + "c" * 64)]),
            "heading_repeated_as_body": (f"{units[0].value}\n{units[0].value}\n{units[1].value}", [units[0], units[1]]),
            "reordered_values": (f"{units[1].value}\n{units[0].value}", [units[0], units[1]]),
            "omitted_value": (units[0].value, [units[0], units[1]]),
        }
        for name, (raw, expected) in cases.items():
            with self.subTest(name=name):
                pdf_units, failures = validation._pdf_body_units(_pdf_raw_extraction(raw), expected)
                self.assertFalse(pdf_units)
                self.assertTrue(failures)

    def test_adapter_enforces_validator_before_descriptors(self) -> None:
        model = _model()
        spec = _spec(model)
        digest = hashlib.sha256(json.dumps(spec, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")).hexdigest()
        with tempfile.TemporaryDirectory(prefix="stage78b2b1-adapter-") as directory:
            root = Path(directory)
            request = root / "request.json"
            output = root / "out"
            result_path = root / "result.json"
            output.mkdir()
            request.write_text(json.dumps({"specification": spec, "digest": digest, "governance_qualification": None, "pathway_render_model": model}), encoding="utf-8")
            with patch.object(adapter, "validate_pathway_output_equivalence", return_value=(types.SimpleNamespace(ok=False), None)) as validator:
                with patch.object(sys, "argv", ["report_adapter.py", str(request), str(output), str(result_path)]):
                    with self.assertRaises(adapter.AdapterFailure) as raised:
                        adapter.main()
        self.assertTrue(validator.called)
        self.assertEqual(raised.exception.phase, "cross_format_equivalence")
        self.assertEqual(raised.exception.code, "equivalence_failed")
        self.assertFalse(result_path.exists())

    def test_adapter_passes_specification_qualification_and_pdf_binding(self) -> None:
        model = _model()
        spec = _spec(model)
        spec["requested_formats"] = ["docx", "html", "pdf"]
        qualification = _qualification()
        digest = hashlib.sha256(json.dumps(spec, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")).hexdigest()
        calls = []

        def fake_equivalence(*args, **kwargs):
            calls.append((args, kwargs))
            return types.SimpleNamespace(ok=True), None

        class FakePdfRenderer:
            def render(self, docx_path, pdf_path, **_kwargs):
                pdf_path.write_bytes(b"%PDF-1.7\nsynthetic\n%%EOF\n")
                return types.SimpleNamespace(renderer_version="fake-libreoffice")

        with tempfile.TemporaryDirectory(prefix="stage78b2b1-adapter-pdf-") as directory:
            root = Path(directory)
            request = root / "request.json"
            output = root / "out"
            result_path = root / "result.json"
            output.mkdir()
            request.write_text(json.dumps({"specification": spec, "digest": digest, "governance_qualification": qualification, "pathway_render_model": model, "governed_job_id": 417, "governed_attempt_count": 1}), encoding="utf-8")
            with patch.object(adapter, "PdfRenderer", return_value=FakePdfRenderer()):
                with patch.object(adapter, "_validate_pdf", return_value={"page_count": 1}):
                    with patch.object(adapter, "validate_pathway_output_equivalence", side_effect=fake_equivalence):
                        with patch.object(sys, "argv", ["report_adapter.py", str(request), str(output), str(result_path)]):
                            adapter.main()
        self.assertEqual(calls[0][0][0], model)
        self.assertEqual(calls[0][1]["specification"], spec)
        self.assertEqual(calls[0][1]["governance_qualification"], qualification)
        self.assertEqual(calls[0][1]["trusted_pdf_conversion"]["conversion_profile"], "internal_pathway_v1")
        self.assertEqual(calls[0][1]["trusted_pdf_conversion"]["converter_version"], "fake-libreoffice")
        self.assertEqual(calls[0][1]["trusted_pdf_conversion"]["conversion_contract"], validation.PATHWAY_TRUSTED_PDF_CONVERSION_CONTRACT)
        self.assertEqual(calls[0][1]["trusted_pdf_conversion"]["job_identity"], "stage77-report-job:417")
        self.assertEqual(calls[0][1]["trusted_pdf_conversion"]["attempt_identity"], "stage77-report-job:417:attempt:1")
        self.assertEqual(calls[0][1]["trusted_pdf_conversion"]["specification_sha256"], digest)
        self.assertEqual(calls[0][1]["trusted_pdf_conversion"]["render_model_identity"], model["render_model_contract"])
        self.assertEqual(calls[0][1]["trusted_pdf_conversion"]["governance_qualification_identity"], qualification["qualification_id"])
        self.assertEqual(calls[0][1]["governed_job_id"], 417)
        self.assertEqual(calls[0][1]["governed_attempt_count"], 1)

    def test_adapter_requires_valid_governed_job_id_for_pathway_pdf(self) -> None:
        model = _model()
        spec = _spec(model)
        spec["requested_formats"] = ["docx", "html", "pdf"]
        digest = hashlib.sha256(json.dumps(spec, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")).hexdigest()
        with tempfile.TemporaryDirectory(prefix="stage78b2b1-adapter-job-id-") as directory:
            root = Path(directory)
            request = root / "request.json"
            output = root / "out"
            result_path = root / "result.json"
            output.mkdir()
            for governed_job_id in (None, True, 0, -1, "417", 417.0):
                with self.subTest(governed_job_id=governed_job_id):
                    payload = {"specification": spec, "digest": digest, "governance_qualification": None, "pathway_render_model": model}
                    if governed_job_id is not None:
                        payload["governed_job_id"] = governed_job_id
                    request.write_text(json.dumps(payload), encoding="utf-8")
                    with patch.object(sys, "argv", ["report_adapter.py", str(request), str(output), str(result_path)]):
                        with self.assertRaises(adapter.AdapterFailure) as raised:
                            adapter.main()
                    self.assertEqual(raised.exception.phase, "input_validation")
                    self.assertEqual(raised.exception.code, "adapter_input_invalid")
                    self.assertFalse(result_path.exists())

    def test_adapter_validator_import_is_mandatory_and_stage75_path_is_unchanged(self) -> None:
        self.assertEqual(Path(adapter._output_validation.__file__).resolve(), (ENGINE / "output_validation.py").resolve())
        self.assertIs(adapter.validate_pathway_output_equivalence, adapter._output_validation.validate_pathway_output_equivalence)
        book = adapter.make_book({
            "title": "Canonical",
            "purpose": "Purpose",
            "specification_schema_version": "stage75.canonical_record_report_specification.v1",
            "report_type": "canonical_record_report",
            "sections": [{"order": 0, "title": "Section", "blocks": []}],
            "selected_documents": [],
            "selected_associations": [],
            "exclusions": [],
            "qualifications": [],
        })
        self.assertEqual(book.title, "Canonical")

    def test_no_ocr_or_tesseract_contract_is_referenced(self) -> None:
        output_source = (ENGINE / "output_validation.py").read_text(encoding="utf-8")
        adapter_source = (ENGINE / "report_adapter.py").read_text(encoding="utf-8")
        combined = output_source + adapter_source
        self.assertNotIn("tesseract", combined.lower())
        self.assertNotIn("ocr", combined.lower())
        self.assertNotIn("pillow", combined.lower())
        self.assertNotIn("def validate_pathway_output_equivalence(*_args, **_kwargs)", adapter_source)

    def test_post_execution_conversion_event_rejects_self_consistent_wrong_job(self) -> None:
        model = _model(); spec = _spec(model); qualification = _qualification()
        with tempfile.TemporaryDirectory(prefix="stage78b2b1-post-event-") as directory:
            root = Path(directory); docx_path = root / "report.docx"; pdf_path = root / "report.pdf"
            _write_framed_docx(docx_path, _units(model, spec, qualification))
            pdf_path.write_bytes(b"%PDF-1.7\nsynthetic\n%%EOF\n")
            event = _trusted_pdf_conversion(model, spec, qualification, docx_path, pdf_path)
            event["job_identity"] = "stage77-report-job:418"
            event["attempt_identity"] = "stage77-report-job:418:attempt:1"
            event["conversion_record_digest"] = validation._pathway_sha256({key: value for key, value in event.items() if key != "conversion_record_digest"})
            with self.assertRaises(ValueError) as raised:
                validation.verify_pathway_pdf_conversion_event(event, docx_path=docx_path, pdf_path=pdf_path, specification=spec, model=model, governance_qualification=qualification, governed_job_id=417, governed_attempt_count=1, retry_of_job_id=None)
        self.assertNotIn("Kind:", str(raised.exception))


if __name__ == "__main__":
    unittest.main()
